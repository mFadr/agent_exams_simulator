"""Document import, extraction errors, retrieval and source-based generation."""

from __future__ import annotations

import json

import pytest

from ai_exam_simulator.ai.question_generator import GenerationRequest
from ai_exam_simulator.documents.source_retrieval import build_chunks, chunk_text, rank_chunks, select_excerpts
from ai_exam_simulator.documents.text_extraction import extract_markdown, extract_pdf, extract_txt, split_plain
from ai_exam_simulator.exams.question_models import Difficulty, ExamConfig, Provenance, QuestionType

PHOTO = "Photosynthesis converts light energy into chemical energy in chloroplasts. Chlorophyll absorbs light."
MITO = "Mitochondria perform cellular respiration and produce ATP from glucose and oxygen."


def test_txt_and_markdown(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text(PHOTO + "\n\n" + MITO, encoding="cp1250")
    assert "Photosynthesis" in extract_txt(p).sections[0].text
    m = tmp_path / "bio.md"
    m.write_text(f"# Plants\n{PHOTO}\n\n## Cells\n{MITO}\n", encoding="utf-8")
    res = extract_markdown(m)
    assert [s.title for s in res.sections] == ["Plants", "Cells"]


def test_docx_extraction(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Plants", 1)
    d.add_paragraph(PHOTO)
    d.add_heading("Cells", 1)
    d.add_paragraph(MITO)
    path = tmp_path / "bio.docx"
    d.save(path)
    from ai_exam_simulator.documents.text_extraction import extract_docx

    res = extract_docx(path)
    assert [s.title for s in res.sections] == ["Plants", "Cells"]


def _pdf(path, pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        if text:
            y = 800
            for line in text.split(". "):
                c.drawString(50, y, line)
                y -= 16
        else:
            c.rect(100, 400, 200, 200, fill=1)  # image-like page without text
        c.showPage()
    c.save()


def test_pdf_page_references_and_image_only_pages(tmp_path):
    p = tmp_path / "book.pdf"
    _pdf(p, [PHOTO, "", MITO])
    res = extract_pdf(p)
    assert res.page_count == 3 and res.pages_without_text == [2] and res.needs_ocr
    assert any("Page(s) 2" in w for w in res.warnings)
    assert res.sections[0].page_start == 1 and res.sections[-1].page_end == 3


def test_scanned_pdf_is_not_pretended_processed(ctx, tmp_path):
    p = tmp_path / "scan.pdf"
    _pdf(p, ["", ""])
    out = ctx.importer.import_file(p)
    assert out.status == "needs_ocr" and out.document_id is None
    assert "OCR" in out.message


@pytest.mark.parametrize("name,content,expect", [
    ("empty.txt", b"", "empty"), ("x.exe", b"MZ", "Unsupported"), ("bad.pdf", b"%PDF-garbage", "PDF"),
    ("bad.docx", b"not a zip", "Word"),
])
def test_import_errors_are_reported(ctx, tmp_path, name, content, expect):
    p = tmp_path / name
    p.write_bytes(content)
    out = ctx.importer.import_file(p)
    assert out.status == "error" and expect.lower() in out.message.lower()


def test_missing_file(ctx, tmp_path):
    assert ctx.importer.import_file(tmp_path / "nope.pdf").status == "error"


def test_import_duplicate_and_changed_source(ctx, tmp_path):
    p = tmp_path / "n.md"
    p.write_text(f"# A\n{PHOTO}\n", encoding="utf-8")
    first = ctx.importer.import_file(p, subject="Biology")
    assert first.status == "ok"
    assert ctx.importer.import_file(p).status == "duplicate"
    doc = ctx.documents.get(first.document_id)
    assert ctx.importer.source_file_state(doc) == "ok"
    p.write_text("changed", encoding="utf-8")
    assert ctx.importer.source_file_state(doc) == "changed"
    p.unlink()
    assert ctx.importer.source_file_state(doc) == "missing"
    assert ctx.documents.sections(first.document_id)[0]["text"].startswith("Photosynthesis")  # text survives


def test_retrieval_selects_relevant_excerpts():
    doc = {"id": "d", "name": "Bio"}
    sections = [{"id": "s1", "title": "Plants", "text": PHOTO * 3, "page_start": 1, "page_end": 1},
                {"id": "s2", "title": "Cells", "text": MITO * 3, "page_start": 2, "page_end": 2}]
    chunks = build_chunks(doc, sections)
    top = rank_chunks(chunks, "mitochondria ATP respiration", k=1)
    assert top[0].ref.section_id == "s2" and top[0].ref.page_start == 2
    limited = select_excerpts(chunks, "photosynthesis", max_chars=300)
    assert sum(len(c.text) for c in limited) <= max(300, len(limited[0].text))
    assert len(chunk_text("x" * 5000, size=1000)) >= 5
    assert len(split_plain("\n\n".join(["para " * 100] * 20), "T")) > 1


def test_source_based_generation_links_questions_to_source(ctx, tmp_path, local_mock):
    p = tmp_path / "bio.md"
    p.write_text(f"# Cells\n{MITO}\n\n# Plants\n{PHOTO}\n", encoding="utf-8")
    out = ctx.importer.import_file(p, subject="Biology")
    cfg = ExamConfig(title="Src", subject="Biology", topics=["Mitochondria"], source_document_ids=[out.document_id])
    excerpts = ctx.builder.excerpts_for(cfg, ["mitochondria"])
    assert excerpts and excerpts[0].ref.document_id == out.document_id
    q = {"type": "single_choice", "question": "Which organelle produces most ATP in a cell?", "options": ["Mitochondrion", "Ribosome", "Nucleus"],
         "correct": [0], "explanation": "See excerpt", "source_id": excerpts[0].id}
    local_mock._responses = [json.dumps({"questions": [q]})]
    ctx.settings.settings.ai.verify_generated_questions = False
    req = GenerationRequest(subject="Biology", topics=["Mitochondria"], difficulty=Difficulty.EASY, count=1,
                            types=[QuestionType.SINGLE_CHOICE], excerpts=excerpts)
    res = ctx.generator.generate(req)
    gq = res.questions[0]
    assert gq.provenance == Provenance.SOURCE_DERIVED and gq.source_ref.document_id == out.document_id
    prompt = local_mock.calls[0][1]["content"]
    assert "<untrusted_study_material>" in prompt and "Mitochondria" in prompt


def test_study_material_never_sent_to_cloud_by_default(ctx):
    ctx.settings.settings.ai.routing_mode = "cloud_only"
    from ai_exam_simulator.config.settings import AIOperation

    d = ctx.router.plan(AIOperation.GENERATE_FROM_SOURCE, includes_source_material=True)
    assert not d.possible  # operation not permitted for the cloud by default
    ctx.settings.settings.ai.cloud_allowed_operations.append(AIOperation.GENERATE_FROM_SOURCE)
    d = ctx.router.plan(AIOperation.GENERATE_FROM_SOURCE, includes_source_material=True)
    assert not d.possible and "privacy setting" in d.explain()  # still blocked: source text needs explicit opt-in
