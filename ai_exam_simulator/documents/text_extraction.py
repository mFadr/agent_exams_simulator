"""Local text extraction from PDF, DOCX, TXT and Markdown."""

from __future__ import annotations

import logging
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

MIN_CHARS_PER_PAGE = 25  # below this a PDF page is treated as image-only


class ExtractionError(Exception):
    """User-facing extraction failure."""


@dataclass
class ExtractedSection:
    title: str
    text: str
    page_start: int | None = None
    page_end: int | None = None


@dataclass
class ExtractionResult:
    format: str
    sections: list[ExtractedSection] = field(default_factory=list)
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)
    pages_without_text: list[int] = field(default_factory=list)
    needs_ocr: bool = False

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1250", "cp1252", "latin-1"):
        if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            continue
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode the text file.")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_markdown(text: str, default_title: str) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    title = default_title
    buf: list[str] = []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}(#{1,3})\s+(.+?)\s*#*\s*$", line)
        if m:
            if "".join(buf).strip():
                sections.append(ExtractedSection(title, _clean("\n".join(buf))))
            title, buf = m.group(2).strip(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append(ExtractedSection(title, _clean("\n".join(buf))))
    return sections


def split_plain(text: str, default_title: str, target_chars: int = 3000) -> list[ExtractedSection]:
    """Split plain text into ~target_chars chunks on paragraph boundaries."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    sections: list[ExtractedSection] = []
    buf: list[str] = []
    size = 0
    for p in paras:
        if size + len(p) > target_chars and buf:
            sections.append(ExtractedSection(f"{default_title} – part {len(sections) + 1}", "\n\n".join(buf)))
            buf, size = [], 0
        buf.append(p)
        size += len(p)
    if buf:
        title = default_title if not sections else f"{default_title} – part {len(sections) + 1}"
        sections.append(ExtractedSection(title, "\n\n".join(buf)))
    return sections


def extract_txt(path: Path) -> ExtractionResult:
    text = _clean(read_text_file(path))
    if not text:
        raise ExtractionError("The file contains no readable text.")
    return ExtractionResult("txt", split_plain(text, path.stem))


def extract_markdown(path: Path) -> ExtractionResult:
    text = read_text_file(path)
    secs = split_markdown(text, path.stem)
    if not secs:
        raise ExtractionError("The file contains no readable text.")
    return ExtractionResult("md", secs)


def extract_pasted(text: str, title: str = "Pasted text") -> ExtractionResult:
    text = _clean(text)
    if not text:
        raise ExtractionError("No text was pasted.")
    secs = split_markdown(text, title) if re.search(r"^#{1,3}\s", text, re.M) else split_plain(text, title)
    return ExtractionResult("paste", secs)


def extract_docx(path: Path) -> ExtractionResult:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractionError("DOCX support requires the 'python-docx' package.") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the Word document ({type(exc).__name__}). Is it a valid .docx file?") from exc
    sections: list[ExtractedSection] = []
    title = path.stem
    buf: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        txt = para.text.strip()
        if not txt:
            continue
        if style.startswith("heading") or style == "title":
            if buf:
                sections.append(ExtractedSection(title, _clean("\n".join(buf))))
            title, buf = txt, []
        else:
            buf.append(txt)
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            buf.append("\n".join(rows))
    if buf:
        sections.append(ExtractedSection(title, _clean("\n".join(buf))))
    res = ExtractionResult("docx", sections)
    if not sections:
        raise ExtractionError("The Word document contains no readable text (images are not processed).")
    if any(rel.reltype.endswith("/image") for rel in document.part.rels.values()):
        res.warnings.append("The document contains images; their content (diagrams, scanned text) was not processed.")
    return res


def extract_pdf(path: Path, *, use_ocr: bool = False) -> ExtractionResult:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:
        raise ExtractionError("PDF support requires the 'pypdf' package.") from exc
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("The PDF is password-protected.") from exc
        pages = reader.pages
        n = len(pages)
    except PdfReadError as exc:
        raise ExtractionError(f"The PDF file is damaged or unsupported ({exc}).") from exc
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read the PDF ({type(exc).__name__}).") from exc

    res = ExtractionResult("pdf", page_count=n)
    page_texts: list[str] = []
    for i in range(n):
        try:
            t = _clean(pages[i].extract_text() or "")
        except Exception as exc:
            log.warning("PDF page %d extraction failed: %s", i + 1, type(exc).__name__)
            t = ""
        if len(t) < MIN_CHARS_PER_PAGE:
            res.pages_without_text.append(i + 1)
        page_texts.append(t)

    if res.pages_without_text and use_ocr:
        ocr_texts = ocr_pdf_pages(path, res.pages_without_text)
        for pno, txt in ocr_texts.items():
            if txt:
                page_texts[pno - 1] = txt
        recovered = [p for p in res.pages_without_text if ocr_texts.get(p)]
        if recovered:
            res.warnings.append(f"OCR was used for page(s) {_ranges(recovered)}; OCR text may contain recognition errors.")
        res.pages_without_text = [p for p in res.pages_without_text if not ocr_texts.get(p)]

    if res.pages_without_text:
        res.needs_ocr = True
        if len(res.pages_without_text) == n:
            res.warnings.append(
                "No text layer was found - this looks like a scanned/image-only PDF. OCR is required"
                + ("" if ocr_available() else " (install Tesseract OCR + 'pytesseract' to enable it)") + "."
            )
        else:
            res.warnings.append(f"Page(s) {_ranges(res.pages_without_text)} have no extractable text (images/scans) and were skipped.")

    # group pages into ~3000-char sections, keeping page references
    buf: list[str] = []
    start = None
    size = 0
    for i, t in enumerate(page_texts, start=1):
        if not t:
            continue
        if start is None:
            start = i
        buf.append(t)
        size += len(t)
        if size >= 3000:
            res.sections.append(ExtractedSection(f"Pages {start}-{i}" if start != i else f"Page {i}", "\n\n".join(buf), start, i))
            buf, start, size = [], None, 0
    if buf and start is not None:
        last = n
        while last > start and not page_texts[last - 1]:
            last -= 1
        res.sections.append(ExtractedSection(f"Pages {start}-{last}" if start != last else f"Page {start}", "\n\n".join(buf), start, last))
    if not res.sections and not res.needs_ocr:
        raise ExtractionError("The PDF contains no readable text.")
    return res


def _ranges(pages: list[int]) -> str:
    out = []
    for p in sorted(pages):
        if out and p == out[-1][1] + 1:
            out[-1][1] = p
        else:
            out.append([p, p])
    return ", ".join(f"{a}" if a == b else f"{a}-{b}" for a, b in out)


# --------------------------------------------------------------------------- optional OCR


def ocr_available() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return bool(shutil.which("tesseract")) or Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").exists()


def ocr_pdf_pages(path: Path, pages: list[int]) -> dict[int, str]:
    """Render pages with Qt PDF and OCR them with Tesseract. Returns {} if unavailable."""
    if not ocr_available():
        return {}
    try:
        import pytesseract
        from PySide6.QtCore import QSize
        from PySide6.QtGui import QImage
        from PySide6.QtPdf import QPdfDocument
    except ImportError:
        return {}
    exe = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if not shutil.which("tesseract") and exe.exists():
        pytesseract.pytesseract.tesseract_cmd = str(exe)
    doc = QPdfDocument()
    if doc.load(str(path)) != QPdfDocument.Error.None_:
        return {}
    out: dict[int, str] = {}
    import tempfile

    with tempfile.TemporaryDirectory(prefix="aies_ocr_") as tmp:
        for p in pages:
            size = doc.pagePointSize(p - 1)
            img: QImage = doc.render(p - 1, QSize(int(size.width() * 2.5), int(size.height() * 2.5)))
            fn = str(Path(tmp) / f"p{p}.png")
            img.save(fn)
            try:
                out[p] = _clean(pytesseract.image_to_string(fn))
            except Exception as exc:
                log.warning("OCR failed on page %d: %s", p, type(exc).__name__)
                out[p] = ""
    return out


def extract_file(path: Path, *, use_ocr: bool = False) -> ExtractionResult:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path, use_ocr=use_ocr)
    if ext == ".docx":
        return extract_docx(path)
    if ext in (".md", ".markdown"):
        return extract_markdown(path)
    if ext == ".txt":
        return extract_txt(path)
    raise ExtractionError(f"Unsupported file type: {ext}")
